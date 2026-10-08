#!/usr/bin/env python3
# Build DOCX, combined Markdown, and a static GitHub Pages site from lexicon term files.
from pathlib import Path
from html import escape
import argparse, re, shutil, yaml
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
CATEGORY_ORDER=[
 ('human-roles','Human Roles'),('ai-roles','AI Roles and Identities'),('relationships','Relationship and Connection Terms'),
 ('interaction','Interaction Modes'),('creativity','Creativity and Innovation'),('workflow','Productivity and Workflow'),
 ('culture','Culture and Humor'),('advanced-states','Advanced Collaboration States'),('future-culture','Networks and Future Culture')]

def parse_entry(path):
 text=path.read_text(encoding='utf-8'); parts=text.split('---',2)
 if len(parts)<3: raise ValueError(f'Missing YAML front matter: {path}')
 meta=yaml.safe_load(parts[1]) or {}; body=parts[2]
 def section(name):
  m=re.search(rf'^## {re.escape(name)}\s*\n(.*?)(?=^## |\Z)',body,re.M|re.S)
  return m.group(1).strip() if m else ''
 neighbors=[]
 for line in section('Semantic Neighbors').splitlines():
  if line.startswith('- '):
   clean=re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',line[2:])
   clean=clean.replace('**','').strip(); neighbors.append(clean)
 return {'path':path,'slug':path.stem,'term':meta.get('term',path.stem),'category':meta.get('category','other'),
  'pos':meta.get('part_of_speech',''),'definition':section('Definition'),'example':section('In use').strip('“”"'),
  'note':section('Usage note'),'neighbors':neighbors,'history':section('History')}

def load_entries():
 entries=[parse_entry(p) for p in sorted((ROOT/'lexicon/terms').glob('*.md'))]
 if not entries: raise SystemExit('No term files found in lexicon/terms.')
 return entries

def add_field(par,instruction,placeholder='Update field in Word'):
 run=par.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
 instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=instruction
 sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); txt=OxmlElement('w:t'); txt.text=placeholder
 end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); run._r.extend([begin,instr,sep,txt,end])

def build_docx(entries,out):
 doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
 styles=doc.styles; styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10.5)
 for name,size,color in [('Title',29,'315C52'),('Heading 1',19,'315C52'),('Heading 2',14,'8A692D')]:
  styles[name].font.name='Aptos Display';styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor.from_string(color)
 if 'Lexicon Term' not in styles:
  st=styles.add_style('Lexicon Term',WD_STYLE_TYPE.PARAGRAPH);st.font.name='Aptos Display';st.font.size=Pt(14);st.font.bold=True;st.font.color.rgb=RGBColor.from_string('315C52');st.paragraph_format.space_before=Pt(12);st.paragraph_format.keep_with_next=True
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('HUMAN / AI');r.bold=True;r.font.size=Pt(18);r.font.color.rgb=RGBColor.from_string('8A692D')
 p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('Relationship Lexicon')
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('COMMUNITY EDITION · GENERATED FROM THE REPOSITORY');r.bold=True;r.font.color.rgb=RGBColor.from_string('315C52')
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(24);p.add_run('A vocabulary for collaboration, creativity, and rapport between humans and AI systems, without pretending AI is human.').italic=True
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(120);p.add_run(f'{len(entries)} ACCEPTED TERMS · OPEN TO USE, REMIX, REVISE, AND EXPAND').bold=True
 doc.add_page_break();doc.add_heading('About this edition',1)
 doc.add_paragraph('This booklet is generated automatically from the accepted Markdown entries in the Human/AI Relationship Lexicon repository. The repository remains the source of truth. Community proposals mature in the Thinkgarden before human harvest review and possible admission to the Core Lexicon.')
 doc.add_heading('Contents',1);add_field(doc.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u')
 doc.add_paragraph('In Microsoft Word, right-click the contents field and choose Update Field if page numbers are not visible.')
 for cat,title in CATEGORY_ORDER:
  group=sorted([e for e in entries if e['category']==cat],key=lambda e:e['term'].casefold())
  if not group: continue
  doc.add_page_break();doc.add_heading(title,1)
  for e in group:
   doc.add_paragraph(e['term'],style='Lexicon Term')
   p=doc.add_paragraph();p.add_run((e['pos'] or 'term')+' · ').bold=True;p.add_run(e['definition'])
   if e['example']:
    p=doc.add_paragraph();p.add_run('In use: ').bold=True;p.add_run('“'+e['example']+'”')
   if e['note']:
    p=doc.add_paragraph();p.add_run('Usage note: ').bold=True;p.add_run(e['note'])
   if e['neighbors']:
    p=doc.add_paragraph();p.add_run('Semantic Neighbors: ').bold=True;p.add_run(' • '.join(e['neighbors']))
 doc.add_page_break();doc.add_heading('The Thinkgarden',1)
 doc.add_paragraph('New terms begin as seedlings. They may germinate, grow, and become ready for harvest through discussion, testing, voting, and human review. Terms that are not adopted are composted into the archive rather than erased.')
 for s in doc.sections:
  p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('Human/AI Relationship Lexicon · ');add_field(p,'PAGE','1')
 out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)

def build_markdown(entries,out):
 lines=['# Human/AI Relationship Lexicon','','> Generated from the repository term files.','']
 for cat,title in CATEGORY_ORDER:
  group=sorted([e for e in entries if e['category']==cat],key=lambda e:e['term'].casefold())
  if not group: continue
  lines += [f'## {title}','']
  for e in group:
   lines += [f'### {e["term"]}','',f'**{e["pos"]}.** {e["definition"]}','']
   if e['example']: lines += [f'**In use:** “{e["example"]}”','']
   if e['note']: lines += [f'**Usage note:** {e["note"]}','']
   if e['neighbors']: lines += [f'**Semantic Neighbors:** {"; ".join(e["neighbors"])}','']
 out.parent.mkdir(parents=True,exist_ok=True);out.write_text('\n'.join(lines),encoding='utf-8')

def page_shell(title,body,root=''):
 return f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(title)}</title><link rel="stylesheet" href="{root}assets/style.css"></head><body><header><a class="brand" href="{root}index.html">Human/AI Relationship Lexicon</a><nav><a href="{root}terms/index.html">Core Lexicon</a><a href="{root}thinkgarden.html">Thinkgarden</a><a href="{root}downloads.html">Downloads</a><a href="{root}contribute.html">Contribute</a></nav></header><main>{body}</main><footer>Open vocabulary for human/AI collaboration · Community Edition</footer></body></html>'''

def build_site(entries,site):
 if site.exists(): shutil.rmtree(site)
 (site/'assets').mkdir(parents=True);(site/'terms').mkdir()
 css='''
:root{--ink:#173b35;--green:#315c52;--mint:#eaf4e7;--gold:#d6ad4f;--paper:#fbfcf7;--muted:#65736f}*{box-sizing:border-box}body{margin:0;background:var(--paper);color:#24332f;font:17px/1.65 system-ui,-apple-system,Segoe UI,sans-serif}header{position:sticky;top:0;z-index:5;display:flex;gap:2rem;justify-content:space-between;align-items:center;padding:1rem max(5vw,1rem);background:#fff;border-bottom:1px solid #dce5df}a{color:var(--green)}nav{display:flex;gap:1rem;flex-wrap:wrap}.brand{font-weight:800;text-decoration:none}main{max-width:1050px;margin:auto;padding:3rem max(4vw,1rem)}.hero{padding:4rem 0}.eyebrow{color:#8a692d;font-weight:800;letter-spacing:.12em}.hero h1{font-size:clamp(2.6rem,7vw,5.8rem);line-height:.98;color:var(--ink);margin:.2em 0}.lead{font-size:1.25rem;max-width:760px}.button{display:inline-block;background:var(--green);color:white;text-decoration:none;padding:.8rem 1rem;border-radius:.7rem;margin:.3rem}.button.secondary{background:var(--gold);color:#21342f}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:1rem}.card{background:white;border:1px solid #dce5df;border-radius:1rem;padding:1.2rem;box-shadow:0 8px 25px #315c5212}.term-list{columns:3 240px}.term-list a{display:block;padding:.3rem 0}.term h1{font-size:3rem;color:var(--ink)}.meta{color:var(--muted);font-weight:700}.example{background:var(--mint);border-left:5px solid var(--gold);padding:1rem}.neighbors li{margin:.4rem 0}input{width:100%;padding:.9rem;border:1px solid #b9c8c0;border-radius:.6rem;font-size:1rem}footer{padding:2rem;text-align:center;color:var(--muted)}@media(max-width:700px){header{position:static;display:block}nav{margin-top:.7rem}.hero{padding:2rem 0}}
'''
 (site/'assets/style.css').write_text(css,encoding='utf-8')
 cats=''.join(f'<article class="card"><h3>{escape(title)}</h3><p>{sum(1 for e in entries if e["category"]==cat)} accepted terms</p><a href="terms/index.html#{cat}">Explore category →</a></article>' for cat,title in CATEGORY_ORDER)
 body=f'''<section class="hero"><div class="eyebrow">OPEN-SOURCE LIVING LANGUAGE PROJECT</div><h1>Words for working with AI.</h1><p class="lead">A community-built vocabulary for collaboration, creativity, workflow, and rapport between humans and AI systems, without pretending AI is human.</p><a class="button" href="terms/index.html">Browse {len(entries)} terms</a><a class="button secondary" href="thinkgarden.html">Visit the Thinkgarden</a></section><section><h2>Explore the lexicon</h2><div class="grid">{cats}</div></section>'''
 (site/'index.html').write_text(page_shell('Human/AI Relationship Lexicon',body),encoding='utf-8')
 index=['<h1>Core Lexicon</h1><p>Accepted terms, grouped by category. Use the search box to filter the list.</p><input id="q" placeholder="Search terms and definitions…" oninput="filterTerms()">']
 for cat,title in CATEGORY_ORDER:
  group=sorted([e for e in entries if e['category']==cat],key=lambda e:e['term'].casefold())
  index.append(f'<section id="{cat}"><h2>{escape(title)}</h2><div class="term-list">')
  for e in group:index.append(f'<a class="term-link" data-search="{escape((e["term"]+" "+e["definition"]).lower())}" href="{e["slug"]}.html">{escape(e["term"])}</a>')
  index.append('</div></section>')
 index.append('''<script>function filterTerms(){const q=document.getElementById('q').value.toLowerCase();document.querySelectorAll('.term-link').forEach(x=>x.style.display=x.dataset.search.includes(q)?'block':'none')}</script>''')
 (site/'terms/index.html').write_text(page_shell('Core Lexicon',''.join(index),'../'),encoding='utf-8')
 for e in entries:
  lis=''.join(f'<li>{escape(n)}</li>' for n in e['neighbors']) or '<li>No neighbors recorded yet.</li>'
  body=f'''<article class="term"><p class="eyebrow">{escape(e['category'].replace('-',' ').upper())}</p><h1>{escape(e['term'])}</h1><p class="meta">{escape(e['pos'])}</p><h2>Definition</h2><p>{escape(e['definition'])}</p><h2>In use</h2><p class="example">“{escape(e['example'])}”</p><h2>Usage note</h2><p>{escape(e['note'])}</p><h2>Semantic Neighbors</h2><ul class="neighbors">{lis}</ul><p><a href="index.html">← Back to Core Lexicon</a></p></article>'''
  (site/f'terms/{e["slug"]}.html').write_text(page_shell(e['term'],body,'../'),encoding='utf-8')
 seedlings=[]
 for p in sorted((ROOT/'thinkgarden/seedlings').glob('*.md')):
  text=p.read_text(encoding='utf-8');parts=text.split('---',2);meta=yaml.safe_load(parts[1]) or {};body=parts[2]
  m=re.search(r'^## Proposed definition\s*\n(.*?)(?=^## |\Z)',body,re.M|re.S);definition=m.group(1).strip() if m else ''
  seedlings.append((meta.get('term',p.stem),meta.get('status','seed'),definition))
 seedcards=''.join(f'<article class="card"><div class="eyebrow">{escape(st.upper())}</div><h2>{escape(t)}</h2><p>{escape(d)}</p></article>' for t,st,d in seedlings)
 (site/'thinkgarden.html').write_text(page_shell('Thinkgarden',f'<h1>The Thinkgarden 🌱</h1><p>The greenhouse where new lexical seedlings are discussed, tested, revised, and voted on before harvest review.</p><div class="grid">{seedcards}</div>'),encoding='utf-8')
 (site/'contribute.html').write_text(page_shell('Contribute','<h1>Plant a lexical seedling</h1><p>Use the repository’s New Term Proposal issue form. Proposals need a definition, example, rationale, usage boundaries, and possible Semantic Neighbors.</p><p><a class="button" href="https://github.com/elkaris11/human-ai-relationship-lexicon/issues/new/choose">Open the proposal forms</a></p>'),encoding='utf-8')
 (site/'downloads.html').write_text(page_shell('Downloads','<h1>Download the Community Edition</h1><p>Generated automatically from the current accepted term files.</p><p><a class="button" href="downloads/Human_AI_Relationship_Lexicon.docx">Word booklet</a><a class="button secondary" href="downloads/Human_AI_Relationship_Lexicon.pdf">PDF booklet</a><a class="button" href="downloads/Human_AI_Relationship_Lexicon.md">Combined Markdown</a></p>'),encoding='utf-8')
 (site/'.nojekyll').write_text('',encoding='utf-8')

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--output',default='build');a=ap.parse_args();out=ROOT/a.output;out.mkdir(parents=True,exist_ok=True)
 entries=load_entries();build_docx(entries,out/'Human_AI_Relationship_Lexicon.docx');build_markdown(entries,out/'Human_AI_Relationship_Lexicon.md');build_site(entries,out/'site')
 downloads=out/'site/downloads';downloads.mkdir(exist_ok=True)
 for name in ['Human_AI_Relationship_Lexicon.docx','Human_AI_Relationship_Lexicon.md']:
  shutil.copy2(out/name,downloads/name)
 print(f'Built {len(entries)} terms into DOCX, Markdown, and website at {out}')
if __name__=='__main__':main()
